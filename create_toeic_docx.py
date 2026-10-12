import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Master questions dataset
QUESTIONS_DATA = {
    # PART 5: 101 - 130
    101: {
        "text": "If ------- are not satisfied with an item, return it for a full refund within 30 days of purchase.",
        "A": "you", "B": "your", "C": "yours", "D": "yourself", "ans": "A"
    },
    102: {
        "text": "The location of next month's online gaming forum is yet to be -------.",
        "A": "concluded", "B": "prevented", "C": "invited", "D": "decided", "ans": "D"
    },
    103: {
        "text": "Guests were ------- with the table decorations for the company banquet.",
        "A": "impressive", "B": "impressed", "C": "impressing", "D": "impressively", "ans": "B"
    },
    104: {
        "text": "The Shubert Company is ------- effective at helping power plants reduce their carbon dioxide emissions.",
        "A": "once", "B": "far", "C": "early", "D": "very", "ans": "D"
    },
    105: {
        "text": "Mr. Hodges ------- that volunteers sign up to assist with the Hannock River cleanup by Friday.",
        "A": "requesting", "B": "to be requested", "C": "requests", "D": "to request", "ans": "C"
    },
    106: {
        "text": "Last year, the Daejeon English News ------- the number of readers by adding a digital subscription option.",
        "A": "needed", "B": "increased", "C": "joined", "D": "asked", "ans": "B"
    },
    107: {
        "text": "Before the updated design can go into -------, it must be approved by management.",
        "A": "product", "B": "producer", "C": "productive", "D": "production", "ans": "D"
    },
    108: {
        "text": "The economic development summit will be held ------- the Xi'an Trade Tower on September 22.",
        "A": "to", "B": "at", "C": "down", "D": "of", "ans": "B"
    },
    109: {
        "text": "Inclement weather was ------- responsible for the low turnout at Saturday's Exton Music Festival.",
        "A": "largely", "B": "large", "C": "largest", "D": "larger", "ans": "A"
    },
    110: {
        "text": "Our most recent survey was sent to clients just last week, ------- it is too soon to send another one.",
        "A": "when", "B": "since", "C": "so", "D": "finally", "ans": "C"
    },
    111: {
        "text": "------- necklace that is shipped from Gillis Designers is given a thorough quality check.",
        "A": "Whenever", "B": "Also", "C": "All", "D": "Each", "ans": "D"
    },
    112: {
        "text": "Ms. Valdez' sales numbers are good ------- for her to be considered for the employee-of-the-month award.",
        "A": "forward", "B": "even", "C": "ahead", "D": "enough", "ans": "D"
    },
    113: {
        "text": "Aaron Park's new book features photographs of homes designed and built by the homeowners -------.",
        "A": "itself", "B": "himself", "C": "themselves", "D": "ourselves", "ans": "C"
    },
    114: {
        "text": "It is ------- for the audience to hold its applause until the speaker has finished.",
        "A": "enthusiastic", "B": "casual", "C": "exclusive", "D": "customary", "ans": "D"
    },
    115: {
        "text": "Despite ------- that Legend Air would perform poorly with the entry of cheaper competition, it posted strong second-quarter earnings.",
        "A": "predicted", "B": "predictable", "C": "predicts", "D": "predictions", "ans": "D"
    },
    116: {
        "text": "Mr. Nigam was ------- retirement when his boss asked him to be the head of security at the new facility.",
        "A": "under", "B": "ahead of", "C": "nearby", "D": "close to", "ans": "D"
    },
    117: {
        "text": "Main Street Restaurant offers a menu of ------- prepared lunch and dinner meals.",
        "A": "thought", "B": "thoughtfulness", "C": "thoughts", "D": "thoughtfully", "ans": "D"
    },
    118: {
        "text": "------- our public relations manager, Ms. Ghazarian has just been appointed vice president of media relations.",
        "A": "Sincerely", "B": "Immediately", "C": "Solely", "D": "Formerly", "ans": "D"
    },
    119: {
        "text": "Videos of Korean pop music have become very popular ------- adolescents worldwide.",
        "A": "including", "B": "whereas", "C": "among", "D": "within", "ans": "C"
    },
    120: {
        "text": "Milante Shoes ------- altered the firm's marketing strategy after a recent economic shift.",
        "A": "quick", "B": "quickest", "C": "quickly", "D": "quicken", "ans": "C"
    },
    121: {
        "text": "The annual report has been posted online, ------- the director's office has not yet received a printed copy.",
        "A": "but", "B": "why", "C": "with", "D": "once", "ans": "A"
    },
    122: {
        "text": "Hasin Fariz turned a study on the ------- effects of sleep into a best-selling book.",
        "A": "favorable", "B": "favor", "C": "favors", "D": "favorably", "ans": "A"
    },
    123: {
        "text": "Wynston Containers is ------- a yearly shutdown of its factory so that it can be evaluated for safety and efficiency.",
        "A": "involving", "B": "participating", "C": "implementing", "D": "producing", "ans": "C"
    },
    124: {
        "text": "The Girard Botanical Archive has almost 300,000 plant -------, all neatly pressed onto archival paper.",
        "A": "authorities", "B": "specimens", "C": "founders", "D": "specifics", "ans": "B"
    },
    125: {
        "text": "Hotels and universities are ------- to recycle their used mattresses through the city's recycling program.",
        "A": "systematic", "B": "eligible", "C": "familiar", "D": "successful", "ans": "B"
    },
    126: {
        "text": "The ------- to review plans to replace the Tronton Bridge will be scheduled soon.",
        "A": "heard", "B": "hears", "C": "hearing", "D": "hear", "ans": "C"
    },
    127: {
        "text": "Kovox Ltd. aims to optimize quality ------- reducing the impact on the environment.",
        "A": "which", "B": "while", "C": "because", "D": "unless", "ans": "B"
    },
    128: {
        "text": "The grocery store ------- vegetables from out of town until local prices went down last month.",
        "A": "is buying", "B": "will be buying", "C": "has been buying", "D": "had been buying", "ans": "D"
    },
    129: {
        "text": "All Hershel Industries employees must have a valid ID card ------- enter the building.",
        "A": "in order to", "B": "as long as", "C": "regarding", "D": "always", "ans": "A"
    },
    130: {
        "text": "------- and cost factored equally in choosing Cantavox as our main supplier.",
        "A": "Reliability", "B": "Allowance", "C": "Dependence", "D": "Estimation", "ans": "A"
    },

    # PART 6: 131 - 146
    131: {
        "text": "",
        "A": "to provide", "B": "provided", "C": "providing", "D": "to be provided", "ans": "A"
    },
    132: {
        "text": "",
        "A": "except", "B": "amid", "C": "near", "D": "for", "ans": "D"
    },
    133: {
        "text": "",
        "A": "Unfortunately, services will be limited.",
        "B": "We thank you for trusting in PGD Bank over these years.",
        "C": "We apologize for any inconvenience this may cause.",
        "D": "Traffic on the boulevard has increased lately.",
        "ans": "C"
    },
    134: {
        "text": "",
        "A": "investments", "B": "hiring", "C": "hours", "D": "interests", "ans": "C"
    },
    135: {
        "text": "",
        "A": "enjoy", "B": "enjoyed", "C": "enjoying", "D": "will enjoy", "ans": "B"
    },
    136: {
        "text": "",
        "A": "revisions", "B": "promotions", "C": "interviews", "D": "receptions", "ans": "C"
    },
    137: {
        "text": "",
        "A": "substance", "B": "substantiate", "C": "substantially", "D": "substantial", "ans": "D"
    },
    138: {
        "text": "",
        "A": "I have four additional questions to ask you.",
        "B": "I would be able to begin during the first week of February.",
        "C": "I am confident I have the potential for this position.",
        "D": "Thank you for the offer of employment.",
        "ans": "B"
    },
    139: {
        "text": "",
        "A": "I can meet you when you arrive.",
        "B": "Scheduling flights can be quite tricky.",
        "C": "I have long admired your work on regional airports.",
        "D": "There are several dining options at the airport.",
        "ans": "C"
    },
    140: {
        "text": "",
        "A": "excel", "B": "excellent", "C": "excellently", "D": "excelled", "ans": "B"
    },
    141: {
        "text": "",
        "A": "join", "B": "pay", "C": "remind", "D": "defend", "ans": "A"
    },
    142: {
        "text": "",
        "A": "to discuss", "B": "discussing", "C": "discuss", "D": "discussed", "ans": "B"
    },
    143: {
        "text": "",
        "A": "run", "B": "has run", "C": "will run", "D": "ran", "ans": "C"
    },
    144: {
        "text": "",
        "A": "movies", "B": "clothing", "C": "food", "D": "books", "ans": "A"
    },
    145: {
        "text": "",
        "A": "political", "B": "popular", "C": "practical", "D": "preliminary", "ans": "B"
    },
    146: {
        "text": "",
        "A": "The awards will be presented by Hunter Johns.",
        "B": "Renovations to the space are nearly complete.",
        "C": "The later offerings were an even bigger success.",
        "D": "Sales are open to the general public on March 3.",
        "ans": "D"
    },

    # PART 7: 147 - 200
    147: {
        "text": "For whom is the advertisement intended?",
        "A": "Recruiters", "B": "Job seekers", "C": "Local business owners", "D": "Current Sedwick Electronics employees", "ans": "B"
    },
    148: {
        "text": "What is stated about Sedwick Electronics?",
        "A": "It is moving its headquarters.",
        "B": "It offers a training program for new employees.",
        "C": "It requires employees to wear uniforms.",
        "D": "It will have more than one location.", "ans": "D"
    },
    149: {
        "text": "What will happen on March 1?",
        "A": "A shipment will be returned.",
        "B": "A new policy will go into effect.",
        "C": "A promotional sale will take place.",
        "D": "A customer survey will be published.", "ans": "B"
    },
    150: {
        "text": "What is NOT a requirement for a complete refund?",
        "A": "The return must be made at the original purchase location.",
        "B": "The return must be made within a certain time frame.",
        "C": "The item must be returned with all its components.",
        "D": "The item must be returned in the original packaging.", "ans": "A"
    },
    151: {
        "text": "Why would people enroll in the course taught by Ms. Sanders?",
        "A": "To practice designing Web sites",
        "B": "To improve their Internet searches",
        "C": "To get tips on creating spreadsheets",
        "D": "To learn how to advertise on the Internet", "ans": "C"
    },
    152: {
        "text": "What is indicated about Ms. Farid?",
        "A": "She also teaches children.",
        "B": "She is Ms. Sanders' supervisor.",
        "C": "She teaches twice a week.",
        "D": "She used to work as a data analyst.", "ans": "C"
    },
    153: {
        "text": "Why did Ms. Witham contact Ms. Ohara?",
        "A": "To review sales figures",
        "B": "To arrange a store visit",
        "C": "To discuss employee performance reviews",
        "D": "To determine the most convenient train to take", "ans": "B"
    },
    154: {
        "text": "At 4:51 P.M., what does Ms. Ohara most likely mean when she writes, \"You should like things here, too\"?",
        "A": "The Tokyo store is being run according to corporate policy.",
        "B": "Ms. Witham will find the athletic shoes she needs.",
        "C": "Ms. Ohara's associate is always punctual.",
        "D": "The Tokyo store is located next to a popular restaurant.", "ans": "A"
    },
    155: {
        "text": "What did Ms. Tulio most likely do?",
        "A": "Make repairs", "B": "Hire a contractor", "C": "Perform an inspection", "D": "Authorize a construction plan", "ans": "C"
    },
    156: {
        "text": "What part of the structure is in most need of maintenance?",
        "A": "The support elements", "B": "The road surface", "C": "The drainage features", "D": "The safety barriers", "ans": "B"
    },
    157: {
        "text": "What is probably true about the Blaine River Drawbridge?",
        "A": "It was not designed for pedestrian use.",
        "B": "It will be closed for the month of October.",
        "C": "It does not have the required signage.",
        "D": "It is the oldest bridge on the Ridgeline Highway.", "ans": "A"
    },
    158: {
        "text": "What is the purpose of the article?",
        "A": "To clarify where subway riders can locate information",
        "B": "To describe improvements at some subway stations",
        "C": "To announce the creation of a new subway line",
        "D": "To explain why subway schedules will be revised", "ans": "B"
    },
    159: {
        "text": "How does Mr. Miller feel about the plans?",
        "A": "He expects the project to fail.",
        "B": "He is concerned about the cost.",
        "C": "He believes the work is unnecessary.",
        "D": "He has been waiting for the changes.", "ans": "D"
    },
    160: {
        "text": "The word \"stages\" in paragraph 5, line 2, is closest in meaning to",
        "A": "steps", "B": "scenes", "C": "train cars", "D": "platforms", "ans": "A"
    },
    161: {
        "text": "What is indicated about Datoric Systems?",
        "A": "It has increased the spending amount allowed for business dinners.",
        "B": "It will adopt new procedures for filing travel expense reports.",
        "C": "It has office locations in several countries.",
        "D": "It plans to hold a company celebration.", "ans": "B"
    },
    162: {
        "text": "Why is the accounting software mentioned?",
        "A": "To highlight a special capability it has",
        "B": "To encourage staff to install it",
        "C": "To help employees log on to it",
        "D": "To point out that it will be replaced", "ans": "A"
    },
    163: {
        "text": "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong?\n\"Following these steps will enable us to quickly issue your reimbursement payment.\"",
        "A": "[1]", "B": "[2]", "C": "[3]", "D": "[4]", "ans": "B"
    },
    164: {
        "text": "What is the main topic of the article?",
        "A": "New trends in marketing athletic footwear",
        "B": "Increased competition in the African clothing market",
        "C": "Recent growth in the African fashion industry",
        "D": "The largest clothing companies in Africa", "ans": "C"
    },
    165: {
        "text": "What is indicated about Abbi Sportscore?",
        "A": "It sells its products only online.",
        "B": "It manufactures its shoes in Nazret.",
        "C": "It will be moving its main offices soon.",
        "D": "It was the first shoe company in Ethiopia.", "ans": "B"
    },
    166: {
        "text": "What is suggested about the Jumjum and Longa Web sites?",
        "A": "They sell only handcrafted goods.",
        "B": "They receive orders from around the world.",
        "C": "They offer free shipping to London and New York.",
        "D": "They are planning to open retail stores.", "ans": "B"
    },
    167: {
        "text": "In which of the positions marked [1], [2], [3], and [4] does the following sentence best belong?\n\"The city boasts four garment factories, with a fifth scheduled to be built this year in nearby Wonji.\"",
        "A": "[1]", "B": "[2]", "C": "[3]", "D": "[4]", "ans": "C"
    },
    168: {
        "text": "Where do the people most likely work?",
        "A": "At a bookstore", "B": "At a public library", "C": "At a television studio", "D": "At a magazine publisher", "ans": "D"
    },
    169: {
        "text": "Why does Mr. Park apologize?",
        "A": "He sent the wrong file.",
        "B": "He used an old e-mail address.",
        "C": "He missed a project deadline.",
        "D": "He lost an important document.", "ans": "A"
    },
    170: {
        "text": "At 10:30 A.M., what does Ms. Riley most likely mean when she writes, \"That's more like it\"?",
        "A": "The budget is more reasonable.",
        "B": "The color looks better.",
        "C": "The story is more interesting.",
        "D": "The schedule is more realistic.", "ans": "B"
    },
    171: {
        "text": "What will Mr. Ojeda do by the afternoon?",
        "A": "Approve a marketing plan",
        "B": "Produce a sample",
        "C": "Repair a printer",
        "D": "Make copies of an agreement", "ans": "B"
    },
    172: {
        "text": "What is the purpose of the letter?",
        "A": "To seek volunteers for an event",
        "B": "To notify a contest finalist",
        "C": "To sell business consultation services",
        "D": "To offer a small-business loan", "ans": "B"
    },
    173: {
        "text": "What does Ms. Chelvi most likely specialize in?",
        "A": "Law", "B": "Technology", "C": "Finance", "D": "Marketing", "ans": "B"
    },
    174: {
        "text": "The word \"illustrates\" in paragraph 3, line 2, is closest in meaning to",
        "A": "represents", "B": "translates", "C": "lightens", "D": "decorates", "ans": "A"
    },
    175: {
        "text": "What is Ms. Chelvi asked to do by March 10?",
        "A": "Update a Web page design",
        "B": "Give a presentation",
        "C": "Sign a consent form",
        "D": "Pay a fee", "ans": "C"
    },
    176: {
        "text": "According to the e-mail, how did Mr. Riedewald receive reward points?",
        "A": "He won an online contest.",
        "B": "He participated in a customer survey.",
        "C": "He spent a certain amount of money.",
        "D": "He returned an item.", "ans": "B"
    },
    177: {
        "text": "In the e-mail, the phrase \"filling out\" in paragraph 1, line 1, is closest in meaning to",
        "A": "emptying", "B": "supplying", "C": "completing", "D": "expanding", "ans": "C"
    },
    178: {
        "text": "How can customers apply their reward points in an MHF retail store?",
        "A": "By entering their account number",
        "B": "By entering their phone number",
        "C": "By scanning a coupon's bar code",
        "D": "By going to the Client Services Department", "ans": "C"
    },
    179: {
        "text": "According to the receipt, what is true about Mr. Riedewald?",
        "A": "He paid for delivery of the items.",
        "B": "He purchased the items in the evening.",
        "C": "He paid over $70 for all items combined.",
        "D": "He purchased only one item at regular price.", "ans": "D"
    },
    180: {
        "text": "What item did Mr. Riedewald most likely purchase using reward credits?",
        "A": "The picture frame",
        "B": "The ceramic mug",
        "C": "The floral blanket",
        "D": "The photo album", "ans": "A"
    },
    181: {
        "text": "What is a purpose of the letter?",
        "A": "To ask Mr. Morello to write a book",
        "B": "To explain an enclosed contract",
        "C": "To notify Mr. Morello of a payment",
        "D": "To describe an updated personnel policy", "ans": "C"
    },
    182: {
        "text": "What was sent in a previous message to Mr. Morello?",
        "A": "Incorrect contact information",
        "B": "Detailed sales numbers",
        "C": "A list of suggested changes",
        "D": "A link to an electronic book", "ans": "B"
    },
    183: {
        "text": "What does Ms. Wicklin mention about Two Swan Press?",
        "A": "It moved to a new location in October.",
        "B": "It has launched a new program for its fifth anniversary.",
        "C": "It has won an industry award.",
        "D": "It has decided to focus on scientific publications.", "ans": "C"
    },
    184: {
        "text": "What is suggested about Mr. Booth?",
        "A": "He is a Two Swan Press author.",
        "B": "He wrote Business in Our Lives.",
        "C": "He is an acquaintance of Mr. Morello.",
        "D": "He has purchased items from Two Swan Press before.", "ans": "A"
    },
    185: {
        "text": "What is indicated about the order?",
        "A": "It has been delayed.",
        "B": "It has not yet been paid.",
        "C": "It contains multiple books.",
        "D": "It includes free shipping.", "ans": "D"
    },
    186: {
        "text": "What does the e-mail indicate about the mentoring program?",
        "A": "It is popular industry-wide.",
        "B": "The number of participants is limited.",
        "C": "It is designed for staff in the sales division.",
        "D": "Participants must attend an orientation meeting.", "ans": "B"
    },
    187: {
        "text": "How will the junior employees most likely be selected?",
        "A": "They will be chosen from a management-training group.",
        "B": "They will undergo competitive interviews.",
        "C": "They will be evaluated by Mr. Wrigley.",
        "D": "They will be recommended by a local business school.", "ans": "C"
    },
    188: {
        "text": "What is suggested about Ms. Drummond?",
        "A": "She has worked at Broadside Electronics for less than eighteen months.",
        "B": "She has just transferred from another department.",
        "C": "She has received a positive annual review.",
        "D": "She has made many successful presentations abroad.", "ans": "A"
    },
    189: {
        "text": "What is most likely true about Ms. Russo?",
        "A": "She is planning to retire soon.",
        "B": "She has international sales experience.",
        "C": "She has mentored many junior employees.",
        "D": "She recently joined the hiring team.", "ans": "B"
    },
    190: {
        "text": "What benefit from the mentoring program have both Ms. Drummond and Ms. Russo enjoyed?",
        "A": "Increased job satisfaction",
        "B": "Quick promotions",
        "C": "Paycheck bonuses",
        "D": "Clearer career goals", "ans": "A"
    },
    191: {
        "text": "What does the Web site indicate about fences taller than three meters?",
        "A": "They are made of plastic.",
        "B": "They are usually prohibited by law.",
        "C": "They require special transportation.",
        "D": "They must be ordered directly from the manufacturer.", "ans": "B"
    },
    192: {
        "text": "According to the form, why does Ms. Carhart need a temporary fence installed?",
        "A": "To mark a property line",
        "B": "To draw attention to an exhibit",
        "C": "To control a crowd at a special event",
        "D": "To limit public access to a work site", "ans": "D"
    },
    193: {
        "text": "What information does Ms. Carhart fail to give about the fence she needs?",
        "A": "The height of the fence to be erected",
        "B": "The dates when the fence is needed",
        "C": "The perimeter of the area to be enclosed",
        "D": "The number of entrances needed", "ans": "C"
    },
    194: {
        "text": "What is implied about Ms. Carhart's fence project?",
        "A": "There will be a charge for delivery.",
        "B": "The project involves work at several job sites.",
        "C": "Extra workers must be hired to install the fence.",
        "D": "Inspectors must first approve the project.", "ans": "A"
    },
    195: {
        "text": "Why does Mr. Montalbo recommend adding a curtain?",
        "A": "It would act as a noise barrier.",
        "B": "It would help keep in dust.",
        "C": "It would serve as a visual screen.",
        "D": "It would improve safety conditions.", "ans": "C"
    },
    196: {
        "text": "Why did Ms. Romero send the e-mail to Mr. Rodrigues Pereira?",
        "A": "To tell him about an upcoming budget cut",
        "B": "To invite him to a concert",
        "C": "To introduce him to his new assistant",
        "D": "To ask him to arrange an event", "ans": "D"
    },
    197: {
        "text": "What does the e-mail imply about Mr. Rodrigues Pereira?",
        "A": "He recently attended a San Jose Starlings game.",
        "B": "He will be leaving in a few weeks to go on vacation.",
        "C": "He is a professional party planner.",
        "D": "He recently began working for Loftgren Consulting.", "ans": "D"
    },
    198: {
        "text": "According to the flyer, what is a benefit of buying tickets as a group?",
        "A": "Reduced ticket prices",
        "B": "Free food",
        "C": "Front-row seating",
        "D": "T-shirts with the team's logo", "ans": "A"
    },
    199: {
        "text": "How many employees does Loftgren Consulting most likely have?",
        "A": "10", "B": "30", "C": "50", "D": "70", "ans": "C"
    },
    200: {
        "text": "On what date could Loftgren Consulting employees attend a game?",
        "A": "August 13", "B": "August 15", "C": "August 19", "D": "August 22", "ans": "B"
    }
}

# Master Passages for Part 6 & Part 7
PART6_PASSAGES = {
    "131-134": {
        "title": "Letter from PGD Bank",
        "text": (
            "Dear PGD Account Holder,\n\n"
            "PGD Bank strives [BLANK: 131] the highest levels of client security and service. "
            "This applies not only to online- and telephone-based services, but also to our brick-and-mortar locations. "
            "Our three branch offices have proudly been a part of the community [BLANK: 132] a combined total of 40 years.\n\n"
            "To assist you even better in the future, our Smithville branch will be temporarily closed for renovations July 8-22. "
            "[BLANK: 133] In the meantime, our other two regional branches in Pine Grove and Bradford will maintain normal business [BLANK: 134]. "
            "We value your feedback and will respond to any concerns that you may have as soon as possible.\n\n"
            "Sincerely,\n"
            "Edwin Chen, Operations Manager\n"
            "PGD Bank"
        )
    },
    "135-138": {
        "title": "E-mail: Cost Accountant Position",
        "text": (
            "To: Samuel Archerson <sarcherson@vona.co.uk>\n"
            "From: James Darrers <jdarrers@sky.co.uk>\n"
            "Date: 10 January\n"
            "Subject: Cost Accountant position\n\n"
            "Dear Mr. Archerson,\n\n"
            "Thank you for taking the time to meet with me today. I [BLANK: 135] our conversation, "
            "and I remain very interested in the position of cost accountant. "
            "I would welcome the opportunity to return for the third and final round of [BLANK: 136].\n\n"
            "I am confident my years of accounting experience would benefit your firm. "
            "As discussed, over the last ten years, I have helped many companies save a [BLANK: 137] amount of money. "
            "I am especially adept at analysing the day-to-day operations of a business and helping to determine more cost-effective methods.\n\n"
            "I checked regarding your question about a potential start date. [BLANK: 138] I hope to hear from you in the near future.\n\n"
            "Sincerely,\n"
            "James Darrers"
        )
    },
    "139-142": {
        "title": "E-mail: Canterbury Airport Project",
        "text": (
            "To: Mason Wu <mwu@wustudios.co.nz>\n"
            "From: Trent Tuiloma <ttuiloma@canterburyairport.co.nz>\n"
            "Subject: Canterbury Airport project\n"
            "Date: Monday, 2 July\n\n"
            "Dear Mr. Wu,\n\n"
            "Thank you for agreeing to consult on the Canterbury Airport redesign project. [BLANK: 139] "
            "As a result, I am particularly eager to hear your ideas about upgrading our main terminal.\n\n"
            "Can we meet this week? There are a number of [BLANK: 140] restaurants near my office. "
            "If you are available this Friday, we could meet at Celia's Cafe on Cumberland Street. "
            "I would also like a few of my colleagues to [BLANK: 141] us. "
            "They would appreciate [BLANK: 142] ways to enhance the airport user's experience.\n\n"
            "I look forward to hearing from you soon.\n\n"
            "Sincerely,\n"
            "Trent Tuiloma\n"
            "Chairman, Canterbury Airport Redesign Team"
        )
    },
    "143-146": {
        "title": "Article: Shiresberry Film Festival",
        "text": (
            "SHIRESBERRY (February 15) - The second annual Shiresberry Film Festival begins on April 18 and [BLANK: 143] for five weeks. "
            "This year's offerings will not be limited to entries from North America and Europe. "
            "We will also be presenting [BLANK: 144] from Asia and South America. "
            "And everyone's favorite feature from last year's festival will be back: "
            "directors and screenwriters will hold question-and-answer sessions after their films' initial screening. "
            "Make sure you do not miss this [BLANK: 145] event. Tickets always sell out quickly. [BLANK: 146] "
            "Shiresberry Film Club members can now purchase priority tickets. "
            "Visit the Shiresberry Theater box office or www.shiresberrytheater.com."
        )
    }
}

PART7_PASSAGES = {
    "147-148": {
        "title": "Advertisement: Sedwick Electronics Hiring Event",
        "text": (
            "Sedwick Electronics Hiring Event\n"
            "March 2, 10 A.M.-5 P.M.\n"
            "22 Myer Street, Hanover, PA 17331\n\n"
            "Sedwick Electronics is opening a new manufacturing facility in Hanover, Pennsylvania, "
            "and we need to fill many positions. We offer a wonderful work environment and great benefits to our employees.\n\n"
            "Come to the event and hear from employees from our Lancaster facility about their experience, "
            "learn about the open positions, and speak with our recruiters. No RSVP is necessary. "
            "Bring copies of your resume."
        )
    },
    "149-150": {
        "title": "Notice: Aguni Plumbing Supply Returns",
        "text": (
            "Aguni Plumbing Supply Returns\n\n"
            "Beginning March 1 at all Aguni Plumbing Supply locations, customers will be able to come to our stores "
            "to return purchases made online. For a complete refund, the return must be made within 30 days of purchase "
            "and must be accompanied by a receipt. In addition, the merchandise must be returned in the original packaging, "
            "and all components must be included. After 30 days, refunds will be limited to in-store credit only. "
            "Defective items may be exchanged for the same item only."
        )
    },
    "151-152": {
        "title": "Information: Springfield Community School Computer Courses",
        "text": (
            "Springfield Community School\n"
            "Computer Courses\n\n"
            "Internet Safety\n"
            "This course teaches students everything they need to navigate the Web safely.\n"
            "Course ID: 249800:01 | Class Time: Tuesday 5:30-7:30 P.M. | Instructor: Patrick McCann | Room: 211\n"
            "Course ID: 249800:02 | Class Time: Saturday 1:00-3:00 P.M. | Instructor: Nora Farid | Room: 166\n\n"
            "Spreadsheet Basics\n"
            "This course teaches the basics of online spreadsheets. Students will learn how to create effective charts for calculating and analyzing data clearly and easily.\n"
            "Course ID: 225810:01 | Class Time: Thursday 5:30-8:30 P.M. | Instructor: Remi Sanders | Room: 118\n"
            "Course ID: 225810:02 | Class Time: Sunday 1:00-4:00 P.M. | Instructor: Nora Farid | Room: 315"
        )
    },
    "153-154": {
        "title": "Text-Message Chain: Store Visit",
        "text": (
            "Sally Witham (4:47 P.M.)\n"
            "Hi Wakiko. I just finished up here at the Kyoto store. I'll be on the train that arrives in Tokyo at 11:35 tomorrow morning. How should I get to your location?\n\n"
            "Wakiko Ohara (4:48 P.M.)\n"
            "I'll have an associate pick you up at the station. How do things look in Kyoto?\n\n"
            "Sally Witham (4:49 P.M.)\n"
            "The Kyoto store is doing a great job. It has everything that we at the home office are looking for. Athletic shoes and sandals are displayed according to specifications, and sales associates are friendly and knowledgeable.\n\n"
            "Wakiko Ohara (4:51 P.M.)\n"
            "You should like things here, too. Do you want to begin your visit after lunch, say at 2:00?\n\n"
            "Sally Witham (4:52 P.M.)\n"
            "Sounds good. See you tomorrow."
        )
    },
    "155-157": {
        "title": "Report: Blaine River Drawbridge Inspection",
        "text": (
            "Structure: Blaine River Drawbridge\n"
            "Location: Ridgeline Highway, KM 147\n"
            "Main span material: Steel girder\n"
            "Owner: State Highway Agency\n"
            "Age of structure: 30 years\n"
            "Report completed by: Vivian Tulio\n"
            "Date: October 17\n\n"
            "Notes:\n"
            "The bridge is overall structurally sound. Inform Department of Transportation about small cracks in asphalt.\n\n"
            "Bridge component | Rating:\n"
            "- Support elements: 4\n"
            "- Towers: 4\n"
            "- Road surface: 3\n"
            "- Drainage features: 4\n"
            "- Safety barriers: 5\n"
            "- Sidewalk or walkway: 6\n\n"
            "Key to ratings:\n"
            "1: Failed; immediate closure required\n"
            "2: Deteriorated; may fail soon\n"
            "3: Shows deterioration but still functions within acceptable parameters\n"
            "4: Shows minor wear\n"
            "5: New condition\n"
            "6: Not applicable"
        )
    },
    "158-160": {
        "title": "Article: Subway Sound to be Upgraded",
        "text": (
            "Subway Sound to be Upgraded\n\n"
            "BOSTON (April 1) - The public address systems at selected subway stations are scheduled to be refurbished, "
            "the Transit Authority announced this week. The systems are used to make announcements to commuters both on the platforms and in the stations.\n\n"
            "Local commuters welcomed the news, although for some it was long overdue.\n\n"
            "\"It can be pretty difficult to understand the announcements at some of the stations I use most frequently,\" said Ian Miller, "
            "who has taken the subway to work nearly every week for the past eighteen years. \"I had heard the reports about it on TV, and all I can say is that it is about time!\"\n\n"
            "Some of the systems currently in use are more than 30 years old. Worn-out speakers, wiring, microphones, and amplifiers will be replaced with new, more reliable devices. "
            "The work should be completed in October and cost more than $11 million.\n\n"
            "Boston's subway system came together in stages over the course of several years. "
            "The foundational component of the system's Green Line first opened on Tremont Street in the late 1890s. It was the first of its kind in the United States."
        )
    },
    "161-163": {
        "title": "E-mail: Expense Report Procedures",
        "text": (
            "Dear Staff,\n\n"
            "I would like to draw your attention to several new procedures regarding business trip expense reports. -[1]-. "
            "Beginning next month, business-related dining receipts must be accompanied by a listing of each dinner attendee. "
            "Also, please make sure that you do not include receipts for any non-work-related items or activities with your report. -[2]-.\n\n"
            "Finally, note that our accounting software will now automatically calculate for you the total to be reimbursed. "
            "You need only to upload images of your receipts for the software to do this.\n\n"
            "I will be happy to respond to your questions. -[3]-. However, I will be flying to Tokyo this Friday to meet clients, "
            "so I will not be checking e-mail that day. -[4]-.\n\n"
            "Sincerely,\n"
            "Selene Hong\n"
            "Assistant Director, Human Resources Department\n"
            "Datoric Systems"
        )
    },
    "164-167": {
        "title": "Article: E-Commerce Opening Doors for African Fashion Industry",
        "text": (
            "E-Commerce Opening Doors for African Fashion Industry\n\n"
            "ADDIS ABABA (6 May) - Africa's role as a consumer of fashion has been on the rise in recent years. "
            "This trend is largely due to the emergence of e-commerce, which provides Africans the opportunity to buy clothing from retailers with no physical presence on the continent.\n\n"
            "Perhaps more importantly, though, the growth of e-commerce is enabling small-scale African designers to also become producers of fashion, "
            "as they showcase their collections to consumers worldwide. African shopping Web sites like Jumjum and Longa are making the work of African designers "
            "available for purchase not just throughout the continent, but also as far away as London and New York. -[1]-.\n\n"
            "\"African designers are finally gaining visibility,\" says Mazaa Absher, founder of Abbi Sportscore, Africa's fastest-growing athletic footwear company. "
            "\"We have always had terrific design and production capacity here on the continent, but it was hard getting it out into the world. Now we are generating more sales online than we are in our stores.\" -[2]-.\n\n"
            "Even as Ms. Absher has transformed her company into an international powerhouse, she continues to highlight the advantages of manufacturing its products in her home city of Nazret. -[3]-. "
            "Africa's strong textile sector and innovative designs combine tradition and wearability, and this formula is allowing companies like hers to set their sights beyond the continent.\n\n"
            "\"As more cities in Ethiopia—and all over Africa—improve their manufacturing capacity, it will become easier to reach the rest of the world,\" says Ms. Absher. -[4]-."
        )
    },
    "168-171": {
        "title": "Text-Message Chain: Cover Design",
        "text": (
            "Gary Park (10:23 A.M.)\n"
            "I e-mailed you the cover design for our September issue a few minutes ago. Did you receive it?\n\n"
            "Jill Riley (10:26 A.M.)\n"
            "Yes, but is this the latest version? I thought we agreed that the background color should be lighter so the article titles are more visible.\n\n"
            "Gary Park (10:28 A.M.)\n"
            "I forgot—sorry about that! I'm just now sending the file with the most recent version.\n\n"
            "Jill Riley (10:30 A.M.)\n"
            "Opening it now ... That's more like it. I'll forward it to Graphics and request a sample printout.\n\n"
            "Jill Riley (10:35 A.M.)\n"
            "Good morning, Mr. Ojeda. Our new cover design is ready. When do you think you'll have a chance to work on it?\n\n"
            "Frank Ojeda (10:38 A.M.)\n"
            "Send it to me now. I'll have a print copy ready for your approval after lunch."
        )
    },
    "172-175": {
        "title": "Letter: Small Business Challenge",
        "text": (
            "8 February\n"
            "Ms. Mala Chelvi\n"
            "60 Jalan Tun Razak\n"
            "54200 Kuala Lumpur\n\n"
            "Dear Ms. Chelvi,\n\n"
            "We are delighted to inform you that you have been nominated as a finalist for the Small Business Challenge competition this year. "
            "Now in its fifth year, this competition is designed to highlight innovative products and services launched by young entrepreneurs. "
            "The Web application that you developed, which provides a means of matching charitable organizations with volunteers, earned one of the top scores from our panel of judges.\n\n"
            "In the next round of the challenge, you will participate in a live presentation about your product before a panel of expert judges. "
            "The three people with the best presentations will receive one-time grants of MYR 10,000 each to invest in their businesses.\n\n"
            "Please go to sbc.org/competition and submit an outline of your presentation, a brief video that clearly illustrates the use of your application, "
            "and a passport-sized photograph of yourself. You will also need to sign a consent form allowing us to use your name and photo, if needed, in promotional materials on our Web site. "
            "The deadline for submission of these materials is 10 March.\n\n"
            "Best regards,\n"
            "Felix Pang\n"
            "Chairperson, Small Business Challenge Committee"
        )
    },
    "176-180": {
        "title": "E-mail and Receipt: McMann Home Furnishings",
        "text": (
            "To: riedewald@parasur.net.sr\n"
            "From: clientservices@mhf.ca\n"
            "Date: April 2, 12:21 P.M.\n"
            "Subject: Your feedback\n\n"
            "Dear Mr. Riedewald,\n\n"
            "Thank you for filling out the McMann Home Furnishings (MHF) survey. To show our appreciation, we have added reward points to your account. "
            "They can be applied to the purchase of products offered online as well as those offered in our retail stores. "
            "Clearance items and those priced $15.00 and above may not be purchased using credits.\n\n"
            "To use your reward points for an online purchase, select the items you would like to purchase and then check out. "
            "At the bottom of the page, select \"Apply credits.\" The value of the applied credits will appear on your order receipt as a special discount.\n\n"
            "If you would prefer to use reward points at one of our retail locations, you can do so by logging in to your account on our Web site. "
            "Go to the My Rewards page, and then select \"Print as a coupon.\" The coupon will have a bar code that can be scanned at the store's checkout counter.\n\n"
            "Sincerely,\n"
            "Client Services, McMann Home Furnishings\n\n"
            "---\n"
            "Online Order #1157\n"
            "McMann Home Furnishings Store\n"
            "March 19, 11:31 A.M.\n\n"
            "Hand-Painted Picture Frame | Quantity: 1 | Price: $10.00 | Special Discount: -$10.00\n"
            "Sailboat Ceramic Mug | Quantity: 4 | Price: $40.00 | Clearance Discount: -$20.00\n"
            "Floral Blanket | Quantity: 1 | Price: $25.00\n"
            "Photo Album | Quantity: 1 | Price: $34.00 | Seasonal Item Discount: -$17.00\n"
            "Item total: $62.00 | Shipping: Free | Total: $62.00"
        )
    },
    "181-185": {
        "title": "Letter and Order: Two Swan Press",
        "text": (
            "Two Swan Press\n"
            "72 Holywell Road, Edinburgh EH8 8PJ\n"
            "4 December\n\n"
            "Mr. Albert Morello\n"
            "17 Peyton Avenue\n"
            "Kingston 5, Jamaica, W.I.\n\n"
            "Dear Mr. Morello:\n\n"
            "Enclosed please find your royalty payment for Understanding Our Oceans. You should have recently received an e-mail that listed the sales figures and the royalties due to you for the print and electronic versions of your book.\n\n"
            "We are proud to announce that Two Swan Press was given the Publisher of the Year Award by the UK Book Industry in October. We thank the authors who have worked with us since our founding five years ago.\n\n"
            "All Two Swan Press authors are entitled to an author discount of 40 percent off any title on our Web site. Simply use the code AUX1417 for your discount.\n\n"
            "If you have any questions at all, please do not hesitate to contact me.\n\n"
            "Kind regards,\n"
            "Sarah Wicklin\n"
            "Encl.\n\n"
            "---\n"
            "https://www.twoswanpress.co.uk/orderconfirmation\n"
            "THANK YOU FOR YOUR ORDER!\n"
            "Special December Offer - free shipping on all orders over £35\n\n"
            "Name: Duncan Booth\n"
            "E-mail: mbooth@silvertech.co.uk\n"
            "Date of purchase: 12 December\n"
            "Ship to: Duncan Booth, 321 Maslin St., Coatbridge ML5 1LZ, Scotland, UK\n\n"
            "1 Business in Our Lives by Elaine Schuyler: £75.00\n"
            "Discount Applied (AUX1417): -£30.00\n"
            "Balance Due: £45.00\n"
            "Paid by Credit Card ****5732\n\n"
            "Items from multiple orders may be combined in the same package. We will notify you when your order has shipped."
        )
    },
    "186-190": {
        "title": "E-mail, Form, and Article: Mentoring Program",
        "text": (
            "To: All Staff\n"
            "From: Personnel Department\n"
            "Date: June 20\n"
            "Subject: Mentoring Program\n"
            "Attachment: Application\n\n"
            "Employees who have been with Broadside Electronics for less than eighteen months are invited to apply to participate in a new mentoring program that will match a maximum of ten junior employees with long-term company veterans. The goal is that junior employees will sharpen corporate skills, better understand company culture, and develop a more focused career path. Mentees will be assigned to a mentor based strictly on their work assignment and professional interests. The pairs will meet at mutually convenient times throughout the year, from three to five hours per month.\n\n"
            "To be considered for participation in this initiative, complete the attached application and return to Mentoring Program Director Tim Wrigley at t.wrigley@broadsideelec.com by July 1. Mr. Wrigley will send notification of his selections by July 15.\n\n"
            "---\n"
            "MENTORING PROGRAM APPLICATION\n"
            "Name: Cora Drummond | Division: Sales | Extension: 144\n"
            "Professional areas of interest: I am most interested in learning about our markets abroad and developing my sales-presentation abilities for these international markets. I am also interested in general career guidance.\n"
            "Best workdays and times for meeting: Any weekday morning except Monday.\n\n"
            "---\n"
            "The Broadside Company Newsletter\n"
            "Mentoring Program Sees Results\n\n"
            "Long-time employee and Vice President of Sales Alena Russo was intrigued when a Personnel Department director approached her about mentoring a less experienced employee under a program that began last year. She is glad to have accepted the assignment. \"After working with Ms. Drummond, I am more satisfied with my own duties, because I know I have helped a professional who is just getting started. I only wish that I had had someone looking out for me in my early years,\" remarked Ms. Russo.\n\n"
            "Ms. Drummond explains that she \"needed pointers on how to make better sales pitches.\" She reports that her sales are up by 20 percent now. She better understands the opportunities Broadside Electronics has to offer and what is required to become a manager. \"Thanks to Ms. Russo, I have been able to define my career goals, and I am a happier person when I arrive to work every day.\"\n\n"
            "New mentorship pairs are now being formed. Interested parties should contact Tim Wrigley in the Personnel Department."
        )
    },
    "191-195": {
        "title": "Web Site, Online Form, and E-mail: Runkle Fencing Service",
        "text": (
            "https://www.runklefencing.co.uk\n"
            "Temporary Fencing Service\n\n"
            "Need a temporary fence installed at your residence, workplace, or special event? Our expert team will deliver and install chain-link fencing on the booked date and time. When the fence is no longer needed, we will take it down and haul it away. Like our residential and commercial fencing, our temporary fences are obtained from the best manufacturers using the finest materials. Our fences have a tidy, sturdy, professional appearance that local residents and businesses appreciate. We also see to all legally required inspections and certifications.\n\n"
            "Ready for a price quote? Contact us today via our Online Quote service. Be sure to provide the following information so we can provide an accurate estimate:\n"
            "1. The preferred height of your fence (Be aware that local regulations commonly allow a maximum height of three metres.)\n"
            "2. The perimeter of the area you need to enclose\n"
            "3. The number of gated entrances needed\n"
            "4. The number of days the fence needs to be up\n\n"
            "---\n"
            "ONLINE QUOTE REQUEST\n"
            "Name: Marguerite Carhart | Phone: (0117) 555-9102\n"
            "Installation Address: 438 Stretford Way, Bristol BS5 7TB\n"
            "E-mail: mcarhart@stockporteventcentre.co.uk\n"
            "Today's date: 8 August\n"
            "Fence Details: I need a temporary three-metre-tall fence installed around the Stockport Event Centre within the next two weeks. This is to prevent the public from entering while we make renovations from 18 to 30 August. Two gates are needed so that workers and vehicles may enter and leave the location.\n\n"
            "---\n"
            "To: mcarhart@stockporteventcentre.co.uk\n"
            "From: hmontalbo@runklefencing.co.uk\n"
            "Date: 9 August\n"
            "Subject: Quote Number 080817\n"
            "Attachment: fencequote_mcarhart\n\n"
            "Dear Ms. Carhart,\n\n"
            "Thank you for your enquiry. Please see the attached estimate for the work you requested. Note that the price of delivery is included at no further charge unless a rush order—one providing less than three weeks' notice—is required. This is a rough estimate based on the information you provided. If you ring us at (0117) 555-2938 and provide us with a missing detail, I can give you a more accurate quote.\n\n"
            "You might also consider including a plastic curtain with your order. This would be wrapped around and fastened to the fence, hiding the construction site from the view of pedestrians. If this interests you, we can include it in the revised quote.\n\n"
            "Best Regards,\n"
            "Howard Montalbo"
        )
    },
    "196-200": {
        "title": "E-mail, Flyer, and Schedule: Company Outing",
        "text": (
            "To: Daniel Rodrigues Pereira\n"
            "From: Livia Romero\n"
            "Subject: Company outing\n"
            "Date: August 5\n\n"
            "Hello Daniel,\n\n"
            "I hope you are settling in well. I'm sure you have had a busy few weeks. Around this time of year, the office manager typically begins arranging our annual company outing. I think we mentioned this during your interview in June. Previously, we have done things like going to a concert and taking a local river cruise. The outing is always great for morale, and everyone looks forward to it.\n\n"
            "This year, I think it would be a good idea to get tickets to a sporting event. I know that many staff members are fans of the San Jose Starlings baseball team. It should be an evening game when the team is playing at home. We have a budget of $600.00 this year. Looking at the ticket prices, it seems that will be just enough to get a ticket for every staff member.\n\n"
            "I'm sure Elise can assist you with this; she has often helped organize the outings. Let me know if you have any questions.\n\n"
            "Best,\n"
            "Livia Romero\n"
            "Director of Administration, Loftgren Consulting\n\n"
            "---\n"
            "Plan your next event with the San Jose Starlings!\n"
            "Discounted tickets are available for groups of ten or more. The more tickets you buy, the more you save - perfect for family gatherings, company outings, or charity fund-raisers! Get perks such as free tickets for the organizer, discounts on food, and your group's name displayed on the scoreboard.\n\n"
            "Group Ticket Pricing:\n"
            "- 10 tickets: $130.00\n"
            "- 30 tickets: $360.00\n"
            "- 50 tickets: $550.00\n"
            "- 70 tickets: $700.00\n"
            "Contact grouptickets@sanjosestarlings.com or call 408-555-0101 for more information.\n\n"
            "---\n"
            "San Jose Starlings - August Schedule\n"
            "- August 13 (Sunday, 1:05 P.M.) vs Aspen Monarchs | Home\n"
            "- August 15 (Tuesday, 7:05 P.M.) vs Aspen Monarchs | Home\n"
            "- August 19 (Saturday, 1:05 P.M.) vs Philipsburg Pinstripes | Away\n"
            "- August 22 (Tuesday, 7:05 P.M.) vs Philipsburg Pinstripes | Away\n"
            "Purchase tickets online at www.sanjosestarlings.com/tickets."
        )
    }
}

def generate_docx():
    doc = docx.Document()
    
    # Styles
    title_p = doc.add_paragraph()
    run = title_p.add_run("[TEST]")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(8, 127, 130) # #087F82 Teal
    
    doc.add_paragraph("[TITLE: TOEIC Reading Test 07]")
    doc.add_paragraph("[TIME: 75]")
    doc.add_paragraph("[TOTAL_QUESTIONS: 100]")
    doc.add_paragraph("")
    
    # PART 5
    p5_hdr = doc.add_paragraph()
    r = p5_hdr.add_run("[PART: 5]")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(24, 50, 79)
    doc.add_paragraph("[PART_TITLE: Incomplete Sentences]")
    doc.add_paragraph("[DESCRIPTION: Directions: A word or phrase is missing in each of the sentences below. Four answer choices are given below each sentence. Select the best answer to complete the sentence.]")
    doc.add_paragraph("")
    
    for qnum in range(101, 131):
        q = QUESTIONS_DATA[qnum]
        doc.add_paragraph(f"[QUESTION: {qnum}]")
        doc.add_paragraph(q["text"])
        doc.add_paragraph(f"[A] {q['A']}")
        doc.add_paragraph(f"[B] {q['B']}")
        doc.add_paragraph(f"[C] {q['C']}")
        doc.add_paragraph(f"[D] {q['D']}")
        doc.add_paragraph(f"[ANSWER: {q['ans']}]")
        doc.add_paragraph("")
        
    # PART 6
    p6_hdr = doc.add_paragraph()
    r = p6_hdr.add_run("[PART: 6]")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(24, 50, 79)
    doc.add_paragraph("[PART_TITLE: Text Completion]")
    doc.add_paragraph("[DESCRIPTION: Directions: Read the texts that follow. A word, phrase, or sentence is missing in parts of each text. Four answer choices for each question are given below the text. Select the best answer to complete the text.]")
    doc.add_paragraph("")
    
    part6_ranges = [("131-134", 131, 134), ("135-138", 135, 138), ("139-142", 139, 142), ("143-146", 143, 146)]
    for grp_key, start_q, end_q in part6_ranges:
        grp = PART6_PASSAGES[grp_key]
        doc.add_paragraph(f"[GROUP: {grp_key}]")
        doc.add_paragraph(f"[PASSAGE_TITLE: {grp['title']}]")
        doc.add_paragraph("[PASSAGE]")
        doc.add_paragraph(grp["text"])
        doc.add_paragraph("[/PASSAGE]")
        doc.add_paragraph("")
        
        for qnum in range(start_q, end_q + 1):
            q = QUESTIONS_DATA[qnum]
            doc.add_paragraph(f"[QUESTION: {qnum}]")
            doc.add_paragraph(f"[A] {q['A']}")
            doc.add_paragraph(f"[B] {q['B']}")
            doc.add_paragraph(f"[C] {q['C']}")
            doc.add_paragraph(f"[D] {q['D']}")
            doc.add_paragraph(f"[ANSWER: {q['ans']}]")
            doc.add_paragraph("")

    # PART 7
    p7_hdr = doc.add_paragraph()
    r = p7_hdr.add_run("[PART: 7]")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(24, 50, 79)
    doc.add_paragraph("[PART_TITLE: Reading Comprehension]")
    doc.add_paragraph("[DESCRIPTION: Directions: In this part you will read a selection of texts, such as magazine and newspaper articles, e-mails, and instant messages. Each text or set of texts is followed by several questions. Select the best answer for each question.]")
    doc.add_paragraph("")
    
    part7_groups = [
        ("147-148", 147, 148),
        ("149-150", 149, 150),
        ("151-152", 151, 152),
        ("153-154", 153, 154),
        ("155-157", 155, 157),
        ("158-160", 158, 160),
        ("161-163", 161, 163),
        ("164-167", 164, 167),
        ("168-171", 168, 171),
        ("172-175", 172, 175),
        ("176-180", 176, 180),
        ("181-185", 181, 185),
        ("186-190", 186, 190),
        ("191-195", 191, 195),
        ("196-200", 196, 200),
    ]
    
    for grp_key, start_q, end_q in part7_groups:
        grp = PART7_PASSAGES[grp_key]
        doc.add_paragraph(f"[GROUP: {grp_key}]")
        doc.add_paragraph(f"[PASSAGE_TITLE: {grp['title']}]")
        doc.add_paragraph("[PASSAGE]")
        doc.add_paragraph(grp["text"])
        doc.add_paragraph("[/PASSAGE]")
        doc.add_paragraph("")
        
        for qnum in range(start_q, end_q + 1):
            q = QUESTIONS_DATA[qnum]
            doc.add_paragraph(f"[QUESTION: {qnum}]")
            if q["text"]:
                doc.add_paragraph(q["text"])
            doc.add_paragraph(f"[A] {q['A']}")
            doc.add_paragraph(f"[B] {q['B']}")
            doc.add_paragraph(f"[C] {q['C']}")
            doc.add_paragraph(f"[D] {q['D']}")
            doc.add_paragraph(f"[ANSWER: {q['ans']}]")
            doc.add_paragraph("")
            
    doc.add_paragraph("[END_TEST]")
    
    filename = "DeThiGiuaKy_Reading_v07.docx"
    doc.save(filename)
    print(f"Successfully generated '{filename}' with full tags and 100 questions!")

if __name__ == "__main__":
    generate_docx()
